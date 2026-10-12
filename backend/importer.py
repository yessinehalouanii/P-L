import csv, io, re
from pathlib import Path
from .source_catalog import SOURCE_CATALOG

ALIASES = {
    'actual amount':'actual','actuals':'actual','budget amount':'budget','last year':'prior_year','prior year':'prior_year',
    'rooms available':'available_rooms','available room':'available_rooms','rooms occupied':'occupied_rooms','occupied rooms':'occupied_rooms',
    'rooms sold':'rooms_sold','room revenue':'room_revenue','occupancy':'occupancy_pct','occupancy %':'occupancy_pct',
    'avg daily rate':'adr','average daily rate':'adr','rev par':'revpar','revenue per available room':'revpar',
    'gross sales':'gross_sales','net sales':'net_sales','labour cost':'total_labour_cost','labor cost':'total_labour_cost',
    'actual hours':'actual_hours','scheduled hours':'scheduled_hours','overtime':'overtime_hours','overtime hours':'overtime_hours',
    'employee':'employee_name','dept':'department','cost':'total_cost','rating score':'rating','review score':'rating'
}

def norm(s):
    s = str(s or '').strip().lower().replace('&',' and ')
    return re.sub(r'[^a-z0-9]+',' ',s).strip()

def canonical_map(source_key, headers):
    fields=SOURCE_CATALOG[source_key]['fields']; lookup={norm(f):f for f in fields}
    mapped={}; unmapped=[]
    for h in headers:
        n=norm(h)
        if n in lookup: mapped[h]=lookup[n]
        elif n in ALIASES and ALIASES[n] in fields: mapped[h]=ALIASES[n]
        else:
            # loose containment matching
            match=next((f for k,f in lookup.items() if n and (n==k or n in k or k in n)),None)
            if match: mapped[h]=match
            else: unmapped.append(h)
    return mapped, unmapped

def parse_upload(filename, data, source_key):
    ext=Path(filename).suffix.lower()
    rows=[]
    if ext=='.csv':
        text=data.decode('utf-8-sig', errors='replace')
        reader=csv.DictReader(io.StringIO(text))
        raw=list(reader); headers=reader.fieldnames or []
    elif ext in ('.xlsx','.xlsm'):
        from openpyxl import load_workbook
        wb=load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        ws=wb.active
        values=list(ws.iter_rows(values_only=True))
        if not values: return [],{},[]
        headers=[str(x or '').strip() for x in values[0]]
        raw=[{headers[i]: row[i] if i < len(row) else None for i in range(len(headers))} for row in values[1:] if any(v not in (None,'') for v in row)]
    else:
        raise ValueError('Only CSV and XLSX files are supported in this demo.')
    mapping,unmapped=canonical_map(source_key,headers)
    for item in raw:
        row={}
        for original,canonical in mapping.items(): row[canonical]=item.get(original)
        if any(v not in (None,'') for v in row.values()): rows.append(row)
    return rows,mapping,unmapped
